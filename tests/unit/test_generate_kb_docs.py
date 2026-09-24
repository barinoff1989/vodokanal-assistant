"""Проверки генератора синтетических документов базы знаний.

Главное, что здесь сторожится, — **пометка о синтетичности**. Документ без неё
неотличим от настоящего регламента: он в том же формате, с теми же заголовками
и с настоящими нормативными ссылками. Ошибка стоила бы дорого — ответ абоненту
по выдуманной процедуре со ссылкой на реальный пункт постановления выглядит
убедительнее настоящего.

Дисциплина та же, что с синтетическими кодами типов обращений и неутверждённым
текстом о качестве воды: синтетика, принятая за настоящее, — одна и та же ошибка
в разном материале.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_kb_docs import (  # noqa: E402
    FOOTER_MARK,
    OBSERVED_REFS,
    SYNTHETIC_MARK,
    build_documents,
    write_document,
)

docx = pytest.importorskip("docx", reason="python-docx не установлен")


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory) -> list[Path]:
    """Документы собираются во временный каталог, а не в `kb/`.

    Проверка не должна зависеть от того, запускал ли кто-то генератор, и не
    должна затирать уже собранное.
    """
    root = tmp_path_factory.mktemp("kb")
    return [write_document(spec, root) for spec in build_documents()]


def test_документы_собираются(generated: list[Path]):
    assert generated
    assert all(path.exists() and path.stat().st_size > 0 for path in generated)


def test_в_каждом_документе_пометка_первым_абзацем(generated: list[Path]):
    """До содержания, а не после: читающий не должен добраться до текста,
    не увидев, что документ синтетический."""
    from docx import Document

    for path in generated:
        texts = [p.text for p in Document(str(path)).paragraphs[:4]]
        assert any(SYNTHETIC_MARK in t for t in texts), path.name


def test_пометка_в_колонтитуле_и_свойствах(generated: list[Path]):
    """Первый абзац уходит за кадр на скриншоте и при печати со второй страницы.

    Поэтому пометка продублирована там, где переживёт и то и другое."""
    from docx import Document

    for path in generated:
        document = Document(str(path))
        footer = " ".join(p.text for p in document.sections[0].footer.paragraphs)
        assert FOOTER_MARK in footer, path.name
        assert document.core_properties.category == "СИНТЕТИКА", path.name
        assert SYNTHETIC_MARK in (document.core_properties.comments or ""), path.name


def test_у_документов_есть_структура_ради_которой_они_и_нужны(generated: list[Path]):
    """Смысл этих файлов — проверить, уважает ли нарезка границы разделов.

    Документ без заголовков такой проверки не даёт, и генератор, потерявший
    структуру, стал бы бесполезен молча."""
    from docx import Document

    for path in generated:
        document = Document(str(path))
        headings = [
            p.text for p in document.paragraphs if p.style.name.startswith("Heading")
        ]
        assert len(headings) >= 3, f"{path.name}: заголовков {len(headings)}"


def test_вложенные_уровни_есть_хотя_бы_в_одном_документе(generated: list[Path]):
    """Ровные разделы одного уровня нарезку не проверяют."""
    from docx import Document

    levels = set()
    for path in generated:
        for paragraph in Document(str(path)).paragraphs:
            if paragraph.style.name.startswith("Heading"):
                levels.add(paragraph.style.name)
    assert len(levels) >= 2, f"уровни заголовков: {levels}"


def test_нормативные_ссылки_только_наблюдавшиеся():
    """Выдуманная ссылка на пункт постановления — тот же класс ошибки, что
    недостоверные цифры в чужих черновиках.

    **Что она ловит:** номер пункта, вписанный прямо в текст документа мимо
    перечня `OBSERVED_REFS`. Проверено — такая правка роняет проверку.

    **Чего она не ловит, и это её граница:** пополнение самого `OBSERVED_REFS`.
    Добавь туда выдуманный пункт — и сравнение сдвинется на обеих сторонах, как
    сдвигается заглушка, подставленная на место проверяемого. Перечень охраняется не проверкой, а
    происхождением: в нём только
    пункты, наблюдавшиеся в ответах операторов и публичном FAQ, и добавлять туда
    следует так же — из наблюдений."""
    import re

    allowed = set()
    for ref in OBSERVED_REFS.values():
        allowed.update(re.findall(r"№ \d+|п\. \d+[^\s,)]*|ст\. \d+", ref))

    found: set[str] = set()
    for spec in build_documents():
        for section in spec.sections:
            blob = " ".join(
                (*section.paragraphs, *section.bullets, *(c for r in section.table for c in r))
            )
            found.update(re.findall(r"№ \d+|п\. \d+[^\s,)]*|ст\. \d+", blob))

    unexpected = found - allowed
    assert not unexpected, f"ссылки вне наблюдавшегося перечня: {sorted(unexpected)}"


def test_типы_обращений_взяты_из_справочника_владельца():
    """Не наш перечень: коды синтетические, но названия пришли от владельца."""
    from app.taxonomy import DISPLAY_NAMES, InquiryType

    names = {DISPLAY_NAMES[t] for t in InquiryType}
    titles = [spec.title for spec in build_documents() if "contracts/" in spec.path]
    assert titles
    for title in titles:
        assert any(name in title for name in names), title
