"""Проверки разбора документов Word на фрагменты базы знаний.

Здесь сторожатся три свойства, каждое из которых уже однажды нарушалось при
написании разбора:

* **путь заголовков несёт название документа** — иначе у пяти файлов реестра
  договоров разделы называются дословно одинаково и фрагменты неразличимы;
* **пометка о синтетичности переживает разбор** — иначе выдуманная процедура
  попадёт в ответ абоненту неотличимо от настоящей;
* **таблица не теряет смысл** — строка «Изменение числа проживающих · Заявление ·
  Расчётный период» без имён столбцов не читается.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.kb.documents import SYNTHETIC_CATEGORY, load_directory, load_docx

pytest.importorskip("docx", reason="python-docx не установлен")


@pytest.fixture(scope="module")
def documents(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Свой каталог документов, а не `kb/`: проверка не должна зависеть от
    того, запускал ли кто-то генератор."""
    import sys

    root = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(root / "scripts"))
    from generate_kb_docs import build_documents, write_document

    target = tmp_path_factory.mktemp("docs")
    for spec in build_documents():
        write_document(spec, target)
    return target


def test_документы_разбираются_на_разделы(documents: Path):
    sections = load_directory(documents)
    assert len(sections) > len(list(documents.rglob("*.docx"))), (
        "разделов должно быть больше, чем файлов — иначе нарезки не произошло"
    )


def test_у_каждого_фрагмента_свой_заголовок(documents: Path):
    """**Проверка, которую разбор не прошёл с первого раза.**

    У пяти файлов реестра договоров разделы называются одинаково — «1. Когда
    применяется», «2. Документы от абонента». Пока название документа не входило
    в путь, вектор заголовка у всех пяти совпадал, и фрагменты становились
    неразличимы для поиска.
    """
    sections = load_directory(documents)
    headings = [s.heading for s in sections]
    assert len(set(headings)) == len(headings), "заголовки повторяются"


def test_путь_заголовков_начинается_с_названия_документа(documents: Path):
    sections = load_directory(documents)
    for section in sections:
        assert "→" in section.heading, section.heading
        assert section.heading.split("→")[0].strip(), section.heading


def test_вложенный_раздел_несёт_родителя(documents: Path):
    """«2.1» само по себе не значит ничего, с родителем — значит."""
    sections = load_directory(documents)
    nested = [s for s in sections if s.heading.count("→") >= 2]
    assert nested, "во всех документах нет ни одного вложенного раздела"


def test_пометка_о_синтетичности_доходит_до_фрагмента(documents: Path):
    """Иначе выдуманная процедура попадёт в ответ неотличимо от настоящей."""
    sections = load_directory(documents)
    assert sections
    assert all(s.synthetic for s in sections)


def test_служебная_надпись_в_текст_фрагмента_не_идёт(documents: Path):
    """Предупреждение — не знание. В поиске оно только мешало бы, совпадая со
    всем подряд, потому что стоит в каждом документе."""
    sections = load_directory(documents)
    for section in sections:
        assert "СИНТЕТИЧЕСКИЙ ДОКУМЕНТ" not in section.body, section.chunk_id


def test_таблица_из_двух_столбцов_читается_как_пары(documents: Path):
    """«Срок рассмотрения: 5 рабочих дней», а не «Условие: Срок рассмотрения ·
    Значение: 5 рабочих дней»."""
    sections = load_directory(documents)
    conditions = [s for s in sections if "Условия и сроки" in s.heading]
    assert conditions
    assert any("Срок рассмотрения:" in s.body for s in conditions)
    assert not any("Условие: Срок" in s.body for s in conditions)


def test_таблица_из_трёх_столбцов_сохраняет_имена(documents: Path):
    """Здесь заголовок столбца, наоборот, обязателен: без него непонятно, что
    основание, а что срок."""
    sections = load_directory(documents)
    recalc = [s for s in sections if "3. Перерасчёт" in s.heading]
    assert recalc
    assert any("Основание:" in s.body and "Срок:" in s.body for s in recalc)


def test_опознаватели_воспроизводимы(documents: Path):
    """Повторная индексация не должна менять опознаватели без причины."""
    first = [s.chunk_id for s in load_directory(documents)]
    second = [s.chunk_id for s in load_directory(documents)]
    assert first == second


def test_несинтетический_документ_помечен_не_будет(tmp_path: Path):
    """Пометка читается из свойств файла, а не приписывается всем подряд."""
    from docx import Document

    document = Document()
    document.add_heading("Настоящий регламент", level=0)
    document.add_heading("1. Раздел", level=1)
    document.add_paragraph("Текст раздела.")
    path = tmp_path / "real.docx"
    document.save(str(path))

    sections = load_docx(path)
    assert sections
    assert not any(s.synthetic for s in sections)
    assert SYNTHETIC_CATEGORY not in (sections[0].source_title or "")


# --- тип обращения документа (payload-фильтр, ADR-200) --------------------------- #


def _doc_with_subject(tmp_path: Path, subject: str | None) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Порядок работы", level=0)
    document.add_heading("1. Раздел", level=1)
    document.add_paragraph("Текст раздела.")
    if subject is not None:
        document.core_properties.subject = subject
    path = tmp_path / "doc.docx"
    document.save(str(path))
    return path


def test_тип_обращения_читается_из_свойства_файла(tmp_path: Path):
    sections = load_docx(_doc_with_subject(tmp_path, "meter_sealing"))
    assert {s.inquiry_type for s in sections} == {"meter_sealing"}


def test_документ_без_типа_общий(tmp_path: Path):
    assert {s.inquiry_type for s in load_docx(_doc_with_subject(tmp_path, None))} == {None}


def test_неизвестный_тип_не_становится_типом(tmp_path: Path):
    """Опечатка в свойстве не должна молча выпадать из поиска по типу."""
    sections = load_docx(_doc_with_subject(tmp_path, "meter_sealling"))
    assert {s.inquiry_type for s in sections} == {None}


def test_реестр_договоров_размечен_регламенты_общие(documents: Path):
    by_file: dict[str, set[str | None]] = {}
    for s in load_directory(documents):
        by_file.setdefault(Path(s.source_path).stem, set()).add(s.inquiry_type)
    assert by_file["meter_verification"] == {"meter_verification"}
    assert by_file["instrukciya-nachisleniya"] == {None}
